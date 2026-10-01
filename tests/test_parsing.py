"""解析引擎单测：Office 三件套走 markitdown 的 extras（docx/xlsx/pptx）。

三个附加包是 pyproject 的核心依赖（`markitdown[docx,xlsx,pptx]`）——掉了它们，
Office 文档进知识库必解析失败。这几条就是「别把 extras 掉回去」的锁。
"""

from pathlib import Path

from nnnu.services.parsing.service import parse_document

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def test_docx_parses_to_markdown(tmp_path: Path):
    # markitdown 读 docx 走的是 mammoth（不是 python-docx，那个没装也不该为测试再引），
    # 所以这里手搓最小 docx：zip 里放三份 XML，mammoth 认得。
    import zipfile

    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-'
        'officedocument.wordprocessingml.document.main+xml"/>'
        "</Types>"
    )
    rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
        'relationships/officeDocument" Target="word/document.xml"/>'
        "</Relationships>"
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        "<w:body><w:p><w:r><w:t>傅里叶变换把时域信号分解为频域分量。</w:t></w:r></w:p>"
        "</w:body></w:document>"
    )
    path = tmp_path / "signal.docx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", rels)
        archive.writestr("word/document.xml", document)

    parsed = parse_document(path, DOCX)
    assert parsed.ok and parsed.error is None
    assert "傅里叶变换" in parsed.text
    assert parsed.pages == []  # Office 无页码语义（页码是 PDF 专属）


def test_xlsx_parses_to_markdown(tmp_path: Path):
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet["A1"], sheet["B1"] = "知识点", "掌握度"
    sheet["A2"], sheet["B2"] = "傅里叶变换", 90
    path = tmp_path / "scores.xlsx"
    book.save(str(path))

    parsed = parse_document(path, XLSX)
    assert parsed.ok and parsed.error is None
    assert "傅里叶变换" in parsed.text
    assert "90" in parsed.text


def test_pptx_parses_to_markdown(tmp_path: Path):
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])  # 只有标题的版式
    slide.shapes.title.text = "信号与系统"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    box.text_frame.text = "傅里叶变换把时域信号分解为频域分量。"
    path = tmp_path / "lecture.pptx"
    deck.save(str(path))

    parsed = parse_document(path, PPTX)
    assert parsed.ok and parsed.error is None
    assert "信号与系统" in parsed.text
    assert "傅里叶变换" in parsed.text
